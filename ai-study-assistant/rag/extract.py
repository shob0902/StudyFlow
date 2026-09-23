# Pull text out of an uploaded file, keeping the page and heading it came from.
#
# Those two pieces of metadata are what make honest citations possible: everything cited later
# comes from here, never from the model.
import io
import re
from dataclasses import dataclass
from typing import Any
from utils.helpers import StudyAssistantError, log_error, log_step
PDF = "pdf"
TXT = "txt"
MARKDOWN = "md"
DOCX = "docx"
SUPPORTED = (PDF, TXT, MARKDOWN, DOCX)
EXTENSIONS = {
    "pdf": PDF,
    "txt": TXT,
    "text": TXT,
    "md": MARKDOWN,
    "markdown": MARKDOWN,
    "docx": DOCX,
}
MAX_BYTES = 20 * 1024 * 1024
MIN_CHARACTERS = 40
# A document that cannot be read, with a message worth showing the user.
class DocumentError(StudyAssistantError):
    pass
# A run of text with where it came from.
@dataclass(frozen=True)
class ExtractedBlock:
    text: str
    page: int | None = None
    section: str = ""
# Everything pulled out of one file.
@dataclass(frozen=True)
class ExtractedDocument:
    blocks: list[ExtractedBlock]
    pages: int = 0
    kind: str = TXT
    # The whole document as plain text.
    @property
    def text(self) -> str:
        return "\n\n".join(block.text for block in self.blocks)
    @property
    def characters(self) -> int:
        return sum(len(block.text) for block in self.blocks)
# Which extractor to use for a filename.
def kind_for(filename: str) -> str:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    kind = EXTENSIONS.get(extension)
    if kind is None:
        raise DocumentError(
            f"'{filename}' is not a supported file type. Upload a PDF, TXT, Markdown or DOCX file."
        )
    return kind
# Refuse a file before doing any work on it.
def validate(filename: str, data: bytes) -> str:
    if not data:
        raise DocumentError(f"'{filename}' is empty.")
    if len(data) > MAX_BYTES:
        raise DocumentError(
            f"'{filename}' is larger than {MAX_BYTES // (1024 * 1024)} MB. Please split it up."
        )
    return kind_for(filename)
# Tidy extracted text: collapse runaway whitespace, drop the hyphenation PDFs leave behind.
def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\xa0", " ")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
# A markdown or plain-text heading, if this line is one.
def _heading(line: str) -> str:
    stripped = line.strip()
    if stripped.startswith("#"):
        return stripped.lstrip("#").strip()[:120]
    if 3 <= len(stripped) <= 80 and stripped.isupper():
        return stripped.title()
    return ""
# Split plain text into blocks, tracking the most recent heading as the section.
def _blocks_from_text(text: str) -> list[ExtractedBlock]:
    blocks: list[ExtractedBlock] = []
    section = ""
    buffer: list[str] = []
    def flush() -> None:
        body = clean_text("\n".join(buffer))
        if body:
            blocks.append(ExtractedBlock(text=body, page=None, section=section))
        buffer.clear()
    for line in text.split("\n"):
        heading = _heading(line)
        if heading:
            flush()
            section = heading
            continue
        if not line.strip():
            if buffer:
                flush()
            continue
        buffer.append(line)
    flush()
    return blocks
# Read a PDF page by page, so every chunk knows its real page number.
def _extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise DocumentError("PDF support is not installed on this server. Upload a TXT or DOCX file.")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise DocumentError("That PDF is password protected, so its text cannot be read.")
        blocks: list[ExtractedBlock] = []
        for number, page in enumerate(reader.pages, start=1):
            body = clean_text(page.extract_text() or "")
            if body:
                blocks.append(ExtractedBlock(text=body, page=number, section=""))
        return ExtractedDocument(blocks=blocks, pages=len(reader.pages), kind=PDF)
    except DocumentError:
        raise
    except Exception:
        log_error("Could not read a PDF")
        raise DocumentError("That PDF could not be read. It may be scanned images rather than text.")
# Read a DOCX, using its heading styles as section names.
def _extract_docx(data: bytes) -> ExtractedDocument:
    try:
        import docx
    except ImportError:
        raise DocumentError("DOCX support is not installed on this server. Upload a PDF or TXT file.")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        log_error("Could not read a DOCX")
        raise DocumentError("That Word file could not be read. Try saving it again, or upload a PDF.")
    blocks: list[ExtractedBlock] = []
    section = ""
    buffer: list[str] = []
    def flush() -> None:
        body = clean_text("\n".join(buffer))
        if body:
            blocks.append(ExtractedBlock(text=body, page=None, section=section))
        buffer.clear()
    for paragraph in document.paragraphs:
        style = (paragraph.style.name if paragraph.style else "") or ""
        text = paragraph.text.strip()
        if not text:
            continue
        if style.lower().startswith("heading") or style.lower() == "title":
            flush()
            section = text[:120]
            continue
        buffer.append(text)
    flush()
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        body = clean_text("\n".join(row for row in rows if row.strip(" |")))
        if body:
            blocks.append(ExtractedBlock(text=body, page=None, section=section or "Table"))
    return ExtractedDocument(blocks=blocks, pages=0, kind=DOCX)
# Read a text or markdown file, guessing the encoding gently.
def _extract_text(data: bytes, kind: str) -> ExtractedDocument:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    else:
        raise DocumentError("That file's text could not be decoded.")
    return ExtractedDocument(blocks=_blocks_from_text(text), pages=0, kind=kind)
# Extract a validated file into blocks of text with their origins.
def extract(filename: str, data: bytes) -> ExtractedDocument:
    kind = validate(filename, data)
    if kind == PDF:
        document = _extract_pdf(data)
    elif kind == DOCX:
        document = _extract_docx(data)
    else:
        document = _extract_text(data, kind)
    if document.characters < MIN_CHARACTERS:
        raise DocumentError(
            f"'{filename}' has almost no readable text. If it is a scanned PDF, it needs OCR first."
        )
    log_step("RAG", f"Extracted {document.characters} characters from {filename!r} ({kind})")
    return document
